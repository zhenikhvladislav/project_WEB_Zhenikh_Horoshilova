from flask import Flask, render_template, redirect, send_file, request, flash
from flask_login import LoginManager, login_user, logout_user, login_required
from forms.user import RegisterForm, LoginForm
from data import db_session
from data.users import User
from data.resumes import Resume
from flask_login import current_user
import pdfkit
from docx import Document
import os

app = Flask(__name__)
login_manager = LoginManager()
login_manager.init_app(app)
app.config['SECRET_KEY'] = 'your_secret_key'
app.config['UPLOAD_FOLDER'] = 'static/uploads'

# Создаем папку для загрузок, если ее нет
if not os.path.exists(app.config['UPLOAD_FOLDER']):
    os.makedirs(app.config['UPLOAD_FOLDER'])


def main():
    db_session.global_init("db/users.db")
    app.run(port=8080, host='127.0.0.1', debug=True)


@login_manager.user_loader
def load_user(user_id):
    db_sess = db_session.create_session()
    return db_sess.query(User).get(user_id)


@app.route('/register', methods=['GET', 'POST'])
def register():
    form = RegisterForm()
    if form.validate_on_submit():
        db_sess = db_session.create_session()
        if db_sess.query(User).filter(User.email == form.email.data).first():
            return render_template('register.html', title='Регистрация', form=form,
                                   message="Такой пользователь уже есть")
        user = User(email=form.email.data,
                    password=form.password.data)
        user.set_password(form.password.data)
        db_sess.add(user)
        db_sess.commit()
        login_user(user, remember=True)
        return redirect('/resume')
    return render_template('register.html', title='Регистрация', form=form)


@app.route('/login', methods=['GET', 'POST'])
def login():
    form = LoginForm()
    if form.validate_on_submit():
        db_sess = db_session.create_session()
        user = db_sess.query(User).filter(User.email == form.email.data).first()
        if user and user.check_password(form.password.data):
            login_user(user, remember=True)
            return redirect('/resume')
        return render_template('login.html', title="Неправильный пароль", form=form)
    return render_template('login.html', title="Авторизация", form=form)


@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect('/')


@app.route('/resume')
def resume():
    return render_template('resume.html')


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/profile')
@login_required
def profile():
    return render_template('profile.html', user=current_user)


@app.route('/submit_resume', methods=['POST'])
@login_required
def submit_resume():
    try:
        db_sess = db_session.create_session()

        # Проверяем, есть ли уже резюме у пользователя
        resume = db_sess.query(Resume).filter(Resume.user_id == current_user.id).first()
        if not resume:
            resume = Resume(user_id=current_user.id)

        # Обновляем данные резюме
        resume.fullname = request.form.get('fullname')
        resume.gender = request.form.get('gender')
        resume.age = request.form.get('age')
        resume.city = request.form.get('city')
        resume.phone = request.form.get('phone')
        resume.email = request.form.get('email')
        resume.social = request.form.get('social')
        resume.messenger = request.form.get('messenger')
        resume.position = request.form.get('position')
        resume.salary = request.form.get('salary')
        resume.education = request.form.get('education')
        resume.experience = request.form.get('experience')
        resume.skills = request.form.get('skills')
        resume.additional = request.form.get('additional')
        resume.qualities = request.form.get('qualities')

        if not db_sess.query(Resume).filter(Resume.user_id == current_user.id).first():
            db_sess.add(resume)

        db_sess.commit()
        flash('Резюме успешно сохранено!', 'success')
    except Exception as e:
        flash(f'Ошибка при сохранении резюме: {str(e)}', 'error')

    return redirect('/resume')


@app.route('/download_pdf')
@login_required
def download_pdf():
    try:
        # Укажите правильный путь к wkhtmltopdf
        config = pdfkit.configuration(wkhtmltopdf='/usr/bin/wkhtmltopdf')  # Linux
        # config = pdfkit.configuration(wkhtmltopdf='C:/Program Files/wkhtmltopdf/bin/wkhtmltopdf.exe')  # Windows

        db_sess = db_session.create_session()
        resume = db_sess.query(Resume).filter(Resume.user_id == current_user.id).first()

        if not resume:
            flash('Сначала заполните и сохраните резюме', 'error')
            return redirect('/resume')

        # Настройки для PDF
        options = {
            'encoding': 'UTF-8',
            'enable-local-file-access': None,
            'page-size': 'A4',
            'margin-top': '15mm',
            'margin-right': '15mm',
            'margin-bottom': '15mm',
            'margin-left': '15mm'
        }

        rendered = render_template('resume_pdf.html', resume=resume)
        pdf_path = os.path.join(app.config['UPLOAD_FOLDER'], f'resume_{current_user.id}.pdf')

        # Генерация PDF
        pdfkit.from_string(
            rendered,
            pdf_path,
            configuration=config,
            options=options
        )

        # Проверка существования файла
        if not os.path.exists(pdf_path):
            flash('Файл PDF не был создан', 'error')
            return redirect('/resume')

        return send_file(
            pdf_path,
            as_attachment=True,
            download_name=f'resume_{current_user.email}.pdf',
            mimetype='application/pdf'
        )

    except Exception as e:
        app.logger.error(f'PDF generation error: {str(e)}')
        flash(f'Ошибка при создании PDF: {str(e)}', 'error')
        return redirect('/resume')


@app.route('/download_docx')
@login_required
def download_docx():
    try:
        db_sess = db_session.create_session()
        resume = db_sess.query(Resume).filter(Resume.user_id == current_user.id).first()

        if resume:
            # Новая версия для заполненного резюме
            doc = Document()
            doc.add_heading(f"Резюме — {resume.fullname}", 0)
            doc.add_paragraph(f"Контактная информация: {resume.phone}, {resume.email}")
            doc.add_paragraph(f"Город: {resume.city}")
            doc.add_paragraph(f"Желаемая должность: {resume.position}")
            doc.add_paragraph(f"Желаемая зарплата: {resume.salary}")
            doc.add_heading("Образование", level=1)
            doc.add_paragraph(resume.education)
            doc.add_heading("Опыт работы", level=1)
            doc.add_paragraph(resume.experience)
            doc.add_heading("Навыки", level=1)
            doc.add_paragraph(resume.skills)

            doc_path = os.path.join(app.config['UPLOAD_FOLDER'], f'resume_{current_user.id}.docx')
            doc.save(doc_path)
            return send_file(doc_path, as_attachment=True)
        else:
            # Старая версия для незаполненного резюме
            doc = Document()
            doc.add_heading(f"Резюме — {current_user.email}", 0)
            doc.add_paragraph("Навыки: Python, Flask, SQL")
            doc.add_paragraph("Опыт: 2 года веб-разработки")
            doc_path = 'static/resume.docx'
            doc.save(doc_path)
            return send_file(doc_path, as_attachment=True)
    except Exception as e:
        flash(f'Ошибка при создании Word документа: {str(e)}', 'error')
        return redirect('/resume')


@app.route('/samples')
def samples():
    return render_template('samples.html')


if __name__ == '__main__':
    main()
